import io
import re
import hmac
import zipfile
import unicodedata
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path

import openpyxl
import streamlit as st

st.set_page_config(page_title="Tạo chứng từ | 7 pháp nhân", page_icon="📑", layout="wide")
BASE = Path(__file__).resolve().parent
TEMPLATE = BASE / "mau" / "Mau_TLO.xlsx"


def normalized(value):
    s = unicodedata.normalize("NFD", str(value or "").strip().lower())
    return "".join(c for c in s if unicodedata.category(c) != "Mn").replace("đ", "d")


def display(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%d/%m/%Y")
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def amount(value):
    try:
        return float(str(value).replace(",", "").replace("%", "").strip())
    except (ValueError, TypeError):
        return 0.0


def at(row, i):
    return row[i] if i < len(row) else None


def put(sheet, address, value):
    cell = sheet[address]
    if not isinstance(cell, openpyxl.cell.cell.MergedCell):
        cell.value = value


def parse_companies(content):
    book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = book.active
        rows = list(sheet.iter_rows(values_only=True))
        header_index = next((i for i, row in enumerate(rows) if any(normalized(v) == "mst" for v in row)), None)
        if header_index is None:
            raise ValueError("Không tìm thấy dòng tiêu đề MST trong file thông tin công ty.")
        headers = {normalized(v): i for i, v in enumerate(rows[header_index]) if v is not None}
        def ix(*names):
            return next((headers[normalized(name)] for name in names if normalized(name) in headers), None)
        columns = {"tax": ix("MST"), "name": ix("Tên"), "address": ix("Địa chỉ mới", "Địa chỉ"),
                   "bank": ix("Số TK", "Tài khoản"), "representative": ix("Đại diện")}
        if any(columns[k] is None for k in ("tax", "name", "address", "representative")):
            raise ValueError("File cần có các cột MST, Tên, Địa chỉ mới và Đại diện.")
        companies = OrderedDict()
        last_tax = None
        for row in rows[header_index + 1:]:
            tax = display(at(row, columns["tax"]))
            bank = display(at(row, columns["bank"])) if columns["bank"] is not None else ""
            if tax:
                name = display(at(row, columns["name"]))
                if not name:
                    continue
                companies[tax] = {"tax": tax, "name": name,
                    "address": display(at(row, columns["address"])),
                    "representative": display(at(row, columns["representative"])),
                    "banks": [bank] if bank else []}
                last_tax = tax
            elif last_tax and bank:
                companies[last_tax]["banks"].append(bank)
        if not companies:
            raise ValueError("Không đọc được pháp nhân nào.")
        return companies
    finally:
        book.close()


ALIASES = {
    "series": ["Ký hiệu"], "invoice": ["Số hóa đơn"], "date": ["Ngày lập"],
    "buyer": ["DN mua", "Tên người mua"], "person": ["Người mua"],
    "taxid": ["MST người mua"], "address": ["Địa chỉ người mua"],
    "item": ["Tên hàng hóa", "Tên hàng hoá", "Tên hàng"],
    "unit": ["Đơn vị tính", "ĐVT"], "price": ["Đơn giá"],
    "qty": ["Số lượng"], "amount": ["Tổng tiền chưa thuế", "Thành tiền"],
    "rate": ["Thuế suất"], "tax": ["Tiền thuế chi tiết", "Tiền thuế"],
    "words": ["Tổng tiền có VAT chữ"],
}
REQUIRED = ["series", "invoice", "date", "buyer", "item", "unit", "price", "qty", "amount", "rate", "tax"]


def parse_invoices(content):
    book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        rows = book.active.iter_rows(values_only=True)
        header = next(rows)
        columns = {normalized(h): i for i, h in enumerate(header) if h is not None}
        positions = {k: next((columns[normalized(a)] for a in options if normalized(a) in columns), None)
                     for k, options in ALIASES.items()}
        missing = [ALIASES[k][0] for k in REQUIRED if positions[k] is None]
        if missing:
            raise ValueError("Thiếu cột: " + ", ".join(missing))
        groups = OrderedDict()
        for row in rows:
            invoice = display(at(row, positions["invoice"]))
            if invoice:
                key = (display(at(row, positions["series"])), invoice)
                groups.setdefault(key, []).append({k: at(row, i) if i is not None else None
                                                      for k, i in positions.items()})
        return groups
    finally:
        book.close()


def make_document(items, invoice, company, code, order_date, delivery_date):
    book = openpyxl.load_workbook(TEMPLATE)
    order, receipt = book.worksheets[:2]
    first = items[0]
    buyer = display(first["buyer"] or first["person"])
    buyer_address = display(first["address"])
    buyer_tax = display(first["taxid"])
    # Replace every fixed seller identity field in both sheets; no TLO text is retained.
    put(order, "A6", f"Số: {invoice}/DDH/{code}")
    put(order, "A8", f"Hôm nay, ngày {order_date},")
    put(order, "A10", "Kính gửi: " + company["name"])
    put(order, "A11", "Địa chỉ: " + company["address"])
    put(order, "A12", "Mã số thuế: " + company["tax"])
    put(order, "A13", "BÊN ĐẶT HÀNG: " + buyer)
    put(order, "A15", "Địa chỉ: " + buyer_address)
    put(order, "A17", "Mã số thuế: " + buyer_tax)
    put(order, "A42", "2. Địa điểm giao hàng: " + buyer_address)
    put(order, "A43", "3. Thời gian giao hàng: " + delivery_date)
    put(receipt, "A8", f"Hôm nay, ngày {delivery_date}, chúng tôi gồm có:")
    put(receipt, "A10", "BÊN BÁN (BÊN A): " + company["name"])
    put(receipt, "A11", "Đại diện: " + company["representative"])
    put(receipt, "A12", "Địa chỉ: " + company["address"])
    put(receipt, "A13", "Mã số thuế: " + company["tax"])
    put(receipt, "A15", "BÊN MUA (BÊN B): " + buyer)
    put(receipt, "A17", "Địa chỉ: " + buyer_address)
    put(receipt, "A18", "Mã số thuế: " + buyer_tax)
    # Buyer representative is not present in all invoice exports; avoid keeping sample person's name.
    put(order, "A14", "Đại diện: " + display(first["person"]))
    put(receipt, "A16", "Đại diện: " + display(first["person"]))
    put(order, "A16", "Điện thoại: ")
    for row in range(21, 36):
        for col in "ABCDEF":
            put(order, f"{col}{row}", None)
    for row in range(23, 40):
        for col in "ABCDE":
            put(receipt, f"{col}{row}", None)
    for j, item in enumerate(items):
        for col, value in [("A", j + 1), ("B", item["item"]), ("C", item["unit"]),
                           ("D", item["qty"]), ("E", item["price"]), ("F", amount(item["amount"]))]:
            put(order, f"{col}{21+j}", value)
        for col, value in [("A", j + 1), ("B", item["item"]), ("C", item["unit"]), ("D", item["qty"])]:
            put(receipt, f"{col}{23+j}", value)
    put(order, "F37", sum(amount(x["amount"]) for x in items))
    put(order, "F38", sum(amount(x["tax"]) for x in items))
    put(order, "F39", "=F37+F38")
    rates = {display(x["rate"]) for x in items}
    put(order, "A38", "THUẾ VAT" + (f" ({next(iter(rates))})" if len(rates) == 1 else " (theo hóa đơn)"))
    put(order, "A40", "Tổng thanh toán bằng chữ: " + display(first["words"]))
    order.print_area = "A1:F53"
    receipt.print_area = "A1:H43"
    order.sheet_view.showGridLines = False
    receipt.sheet_view.showGridLines = False
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


st.title("HỆ THỐNG TẠO CHỨNG TỪ · ĐA PHÁP NHÂN")
password = st.secrets.get("APP_PASSWORD", "")
if not password:
    st.error("Chưa cấu hình APP_PASSWORD trong Streamlit → App settings → Secrets.")
    st.stop()
entered = st.text_input("Mật khẩu truy cập nội bộ", type="password")
if not hmac.compare_digest(entered, password):
    st.info("Nhập mật khẩu để tiếp tục.")
    st.stop()

st.caption("Thông tin pháp nhân và hóa đơn được tải lên khi sử dụng, không lưu vào GitHub.")
with st.sidebar:
    st.subheader("Ngày chứng từ")
    order_day = st.date_input("Ngày đặt hàng", date.today())
    delivery_day = st.date_input("Ngày giao nhận", date.today())
    use_invoice_day = st.checkbox("Lấy ngày theo hóa đơn", value=False)
    st.warning("Chỉ tải dữ liệu thật khi công ty chấp thuận xử lý trên Streamlit Cloud.")

company_upload = st.file_uploader("1. Tải THÔNG TIN CÔNG TY.xlsx (chỉ quản trị viên giữ file)", type=["xlsx"], key="company")
if not company_upload:
    st.info("Tải file thông tin 7 pháp nhân để bắt đầu. File này KHÔNG được đưa lên GitHub Public.")
    st.stop()
try:
    companies = parse_companies(company_upload.getvalue())
except Exception as exc:
    st.error(f"Không đọc được danh sách pháp nhân: {exc}")
    st.stop()

codes = {"0109865925": "AZ", "0110011822": "PHUONG_HOANG", "0317344602": "TLO",
         "0110024356": "HUY_HOANG", "0110032702": "MD", "0110102798": "THO", "0109703201": "ELIM"}
tax = st.selectbox("2. Chọn pháp nhân", list(companies), format_func=lambda x: f"{codes.get(x, x)} · {companies[x]['name']}")
company = companies[tax]
code = codes.get(tax, re.sub(r"\W+", "_", company["name"], flags=re.UNICODE)[:25])
st.success(f"Đang tạo chứng từ cho: {company['name']} | MST: {tax}")
with st.expander("Kiểm tra thông tin bên bán"):
    st.write("Địa chỉ:", company["address"])
    st.write("Đại diện:", company["representative"])
    if company["banks"]:
        st.selectbox("Tài khoản ngân hàng (tham khảo, mẫu hiện tại chưa có ô tài khoản)", company["banks"])

invoice_upload = st.file_uploader("3. Tải danh sách hóa đơn Excel", type=["xlsx"], key="invoices")
if not invoice_upload:
    st.info("Hãy tải file hóa đơn để tạo chứng từ.")
    st.stop()
try:
    groups = parse_invoices(invoice_upload.getvalue())
except Exception as exc:
    st.error(f"Không đọc được file hóa đơn: {exc}")
    st.stop()
if not groups:
    st.warning("Không tìm thấy hóa đơn hợp lệ.")
    st.stop()
st.success(f"Đọc được {len(groups)} hóa đơn, {sum(map(len, groups.values()))} dòng hàng.")
labels = {f"{series} · {invoice} ({len(items)} dòng)": (series, invoice)
          for (series, invoice), items in groups.items()}
selected = st.multiselect("4. Chọn hóa đơn", list(labels), default=list(labels))
invalid = []
for label in selected:
    key = labels[label]
    items = groups[key]
    if len(items) > 15:
        invalid.append(f"{key[1]}: {len(items)} dòng (mẫu tối đa 15)")
    elif any(not display(x["item"]) or x["qty"] is None for x in items):
        invalid.append(f"{key[1]}: thiếu tên hàng hoặc số lượng")
for warning in invalid:
    st.warning(warning)
valid = [labels[label] for label in selected if not any(labels[label][1] == w.split(":", 1)[0] for w in invalid)]
st.caption(f"Sẵn sàng tạo {len(valid)} bộ; không đủ điều kiện: {len(selected)-len(valid)} bộ.")
if st.button("5. Tạo và đóng gói chứng từ", type="primary", disabled=not valid):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for series, invoice in valid:
            items = groups[(series, invoice)]
            when_order = display(items[0]["date"]) if use_invoice_day else order_day.strftime("%d/%m/%Y")
            when_delivery = display(items[0]["date"]) if use_invoice_day else delivery_day.strftime("%d/%m/%Y")
            data = make_document(items, invoice, company, code, when_order, when_delivery)
            safe_invoice = re.sub(r"[^a-zA-Z0-9_-]", "_", invoice)
            safe_series = re.sub(r"[^a-zA-Z0-9_-]", "_", series)
            archive.writestr(f"{code}_{safe_series}_{safe_invoice}_DonDH_BBGN.xlsx", data)
    st.session_state["generated"] = {"data": output.getvalue(), "tax": tax,
                                     "signature": (invoice_upload.getvalue(), company_upload.getvalue(), tuple(selected),
                                                   str(order_day), str(delivery_day), use_invoice_day), "count": len(valid)}
result = st.session_state.get("generated")
current_signature = (invoice_upload.getvalue(), company_upload.getvalue(), tuple(selected),
                     str(order_day), str(delivery_day), use_invoice_day)
if result and result["tax"] == tax and result["signature"] == current_signature:
    st.download_button(f"⬇️ Tải ZIP ({result['count']} bộ chứng từ)", result["data"],
                       file_name=f"{code}_chung_tu.zip", mime="application/zip")
    st.caption("Mỗi file Excel có 2 sheet: Đơn đặt hàng và Biên bản giao nhận. Các hóa đơn không hợp lệ không được xuất.")
